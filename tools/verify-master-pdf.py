"""Live OOG-only tests; remove only uniquely identified fixtures and our JWT sessions."""
import base64, io, json, secrets, uuid, urllib.request, urllib.error, time
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as ET
from contextlib import closing
import psycopg
import pymupdf as fitz
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
config=dict(line.split('=',1) for line in (ROOT/'backend/.env').read_text(encoding='utf-8-sig').splitlines() if '=' in line)
url=urlparse(config['DB_URL'].removeprefix('jdbc:'));assert url.path=='/oog'
checks=[];tenants=[];sessions=[];projects=[];fixture_code='T_'+uuid.uuid4().hex[:12]
def check(name,ok):
    assert ok,name
    checks.append(name)
def request(path,method='GET',body=None,token=None):
    headers={'Content-Type':'application/json'}
    if token:headers['Authorization']='Bearer '+token
    try:
        with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8083/api'+path,data=None if body is None else json.dumps(body).encode(),headers=headers,method=method),timeout=45) as r:
            raw=r.read()
            return r.status,json.loads(raw)['data'] if r.headers.get_content_type().endswith('json') else raw
    except urllib.error.HTTPError as e:return e.code,e.read()
def login(code,user,password):
    status,value=request('/auth/login','POST',{'tenantCode':code,'username':user,'password':password})
    check('login '+code,status==200)
    sessions.append(json.loads(base64.urlsafe_b64decode(value['accessToken'].split('.')[1]+'=='))['jti'])
    return value['accessToken']

for attempt in range(45):
    try:
        status,_=request('/me')
        if status==401:break
    except urllib.error.URLError:pass
    time.sleep(1)
else:raise RuntimeError('OOG API did not become ready')
platform=login('OOG','admin',config['OOG_ADMIN_PASSWORD'])
check('reference catalog anonymous denied',request('/equipment/references')[0]==401)
status,reference=request('/equipment/references',token=platform)
check('eight source candidates without refrigerated types',status==200 and {c['code']for c in reference['catalog']['candidates']}=={'20GP','40GP','40HC','45HC','20OT','40OT','20FR','40FR'})
check('reference mass identity and explicit non-approval',all(c['tareKg']+c['payload']==c['maximumGrossKg']for c in reference['catalog']['candidates']) and not reference['catalog']['autoApply'] and not reference['catalog']['carrierApproved'])
check('reference source digest recorded',reference['catalog']['sourceImageSha256']=='3dc7679f129544075ce062b04319bad36c2c863e928ddd872df1304016acba45')
with psycopg.connect(host=url.hostname,port=url.port,dbname='oog',user=config['DB_USERNAME'],password=config['DB_PASSWORD'],connect_timeout=10) as db:
    try:
        registration={'code':fixture_code,'specification':{'name':'VERIFY equipment','length':1000,'width':1000,'height':1000,'payload':1000,'openWidth':False,'openHeight':False,'version':0},'usableLengthRatio':.94,'planningWeightRatio':.9,'maxOverWidth':0,'maxOverHeight':0,'evidence':'Code/API verification fixture only'}
        # Public API must create both the physical specification and its planning rule atomically.
        registration['code']=fixture_code=fixture_code.upper()
        check('anonymous equipment registration denied',request('/equipment','POST',registration)[0]==401)
        check('missing registration evidence rejected',request('/equipment','POST',{**registration,'evidence':''},platform)[0]==400)
        check('closed side allowance rejected',request('/equipment','POST',{**registration,'maxOverWidth':1},platform)[0]==400)
        status,created=request('/equipment','POST',registration,platform)
        check('equipment registration API',status==200 and created['code']==fixture_code and created['version']==0)
        check('duplicate equipment code conflict',request('/equipment','POST',registration,platform)[0]==409)
        db.commit()
        check('default planning rule created atomically',db.execute('select count(*) from oog_carrier_rule where equipment_code=%s and usable_length_ratio=.94',(fixture_code,)).fetchone()[0]==1)
        check('registration source audited',db.execute("select count(*) from core_audit_history where action='EQUIPMENT_CREATED' and detail->>'code'=%s and detail->>'evidence'=%s",(fixture_code,registration['evidence'])).fetchone()[0]==1)
        password=secrets.token_urlsafe(24)
        for suffix in ['A','B']:
            code='M_'+uuid.uuid4().hex[:12].upper()
            status,value=request('/core/tenants','POST',{'tenantCode':code,'tenantName':'VERIFY-MASTER-'+suffix,'adminUsername':'verifyadmin','adminPassword':password},platform)
            check('company fixture '+suffix,status==200);tenants.append((value['tenantId'],code))
        a=login(tenants[0][1],'verifyadmin',password);b=login(tenants[1][1],'verifyadmin',password)
        check('company viewer can compare reference but no values auto applied',request('/equipment/references',token=a)[0]==200 and all(not r['autoApplied']for r in request('/equipment/references',token=a)[1]['comparison']))
        check('company admin registration forbidden',request('/equipment','POST',{**registration,'code':fixture_code+'_X'},a)[0]==403)
        specification={'name':'VERIFY edited','length':2000,'width':1500,'height':1200,'payload':3000,'openWidth':True,'openHeight':True,'version':0}
        path='/equipment/'+fixture_code
        check('equipment anonymous denied',request(path,'PUT',specification)[0]==401)
        check('company admin cannot alter global equipment',request(path,'PUT',specification,a)[0]==403)
        check('zero equipment dimension rejected',request(path,'PUT',{**specification,'length':0},platform)[0]==400)
        missing={k:v for k,v in specification.items() if k!='version'}
        check('equipment version required',request(path,'PUT',missing,platform)[0]==400)
        status,updated=request(path,'PUT',specification,platform)
        check('equipment version and dimensions saved',status==200 and updated['version']==1 and updated['length']==2000)
        opening={'mode':'DOOR','cargoLengthMm':1000,'cargoWidthMm':1000,'cargoHeightMm':1000,'openingAmm':1100,'openingBmm':1100,'clearancePerSideMm':50,'rotationAllowed':False,'evidence':'verified opening fixture'}
        check('anonymous opening check denied',request(path+'/clearance','POST',opening)[0]==401)
        status,screen=request(path+'/clearance','POST',opening,a)
        check('company view user can screen exact opening clearance',status==200 and screen['status']=='RECTANGULAR_OPENING_FITS' and not screen['carrierApproved'])
        check('one mm opening failure',request(path+'/clearance','POST',{**opening,'openingAmm':1099},a)[1]['status']=='RECTANGULAR_OPENING_EXCEEDED')
        check('opening evidence required',request(path+'/clearance','POST',{**opening,'evidence':''},a)[0]==400)
        check('opening missing equipment not found',request('/equipment/NO_SUCH_EQUIPMENT/clearance','POST',opening,a)[0]==404)
        check('stale equipment update conflicts',request(path,'PUT',specification,platform)[0]==409)
        with db.cursor() as cur:
            cur.execute("select detail from core_audit_history where action='EQUIPMENT_UPDATED' and detail->>'equipmentCode'=%s",(fixture_code,))
            audit=cur.fetchone()[0]
            check('equipment audit records before and after',audit['before']['length']==1000 and audit['after']['length']==2000)
        rule={'carrier':'VERIFYLINE','equipmentCode':'40HC','validFrom':'2026-01-01','validUntil':'2026-12-31','usableLengthRatio':.5,'planningWeightRatio':.8,'maxOverWidth':0,'maxOverHeight':0,'reason':'Verification fixture'}
        check('invalid rule ratio rejected',request('/carrier-rules','POST',{**rule,'usableLengthRatio':0},a)[0]==400)
        check('invalid effective range rejected',request('/carrier-rules','POST',{**rule,'validUntil':'2025-12-31'},a)[0]==400)
        status,v1=request('/carrier-rules','POST',rule,a);check('company rule version created',status==200 and v1['version']==1)
        check('overlapping carrier dates conflict',request('/carrier-rules','POST',rule,a)[0]==409)
        status,v2=request('/carrier-rules','POST',{**rule,'validFrom':'2027-01-01','validUntil':None,'usableLengthRatio':.9},a)
        check('next effective rule version created',status==200 and v2['version']==2)
        rows_a=request('/carrier-rules',token=a)[1];rows_b=request('/carrier-rules',token=b)[1]
        check('carrier rule list tenant isolated',len(rows_a)==2 and len(rows_b)==0)
        for operation in ['update oog_carrier_rule_version set reason=reason where id=%s','delete from oog_carrier_rule_version where id=%s']:
            db.execute('savepoint immutable_check')
            rejected=False
            try:db.execute(operation,(v1['id'],))
            except psycopg.Error:rejected=True
            db.execute('rollback to savepoint immutable_check')
            check('rule SQL immutable '+operation.split()[0],rejected)
        status,p=request('/projects','POST',{'name':'VERIFY-MASTER-PDF','customer':'검증 고객'},a);check('project fixture created',status==200);projects.append(p['id']);pid=p['id']
        for item,length in [('한글펌프',7000),('UNASSIGNED',100000)]:
            status,_=request('/projects/'+pid+'/cargos','POST',{'projectId':pid,'itemNo':item,'description':'포장 화물','quantity':1,'length':length,'width':1000,'height':1000,'grossWeight':1000},a)
            check('cargo fixture '+item,status==200)
        query='?carrier=VERIFYLINE&effectiveDate=2026-06-01&allowDefaultFallback=true'
        status,analysis=request('/projects/'+pid+'/analysis'+query,token=a)
        check('selected carrier analysis available',status==200 and any(not c['feasible'] for row in analysis if row['cargo']['itemNo']=='한글펌프' for c in row['candidates'] if c['code']=='40HC'))
        meta=analysis[0]['ruleSelection']['versions']
        check('fallback is explicitly recorded',any(v['usedFallback'] and v['source']=='DEFAULT' for v in meta))
        status,end=request('/projects/'+pid+'/analysis?carrier=VERIFYLINE&effectiveDate=2026-12-31',token=a)
        check('end date is inclusive',status==200 and next(v for v in end[0]['ruleSelection']['versions'] if v['equipmentCode']=='40HC')['version']==1)
        status,future=request('/projects/'+pid+'/analysis?carrier=VERIFYLINE&effectiveDate=2027-01-01',token=a)
        check('next effective version used',status==200 and next(v for v in future[0]['ruleSelection']['versions'] if v['equipmentCode']=='40HC')['version']==2)
        check('future carrier rule changes feasibility',any(c['feasible'] for row in future if row['cargo']['itemNo']=='한글펌프' for c in row['candidates'] if c['code']=='40HC'))
        status,raw=request('/projects/'+pid+'/export'+query,token=a)
        with closing(openpyxl.load_workbook(io.BytesIO(raw))) as book:
            values=list(book['OOG Analysis'].values)
            check('analysis Excel applies selected carrier',status==200 and 'OL' in next(row[8] for row in values[1:] if row[0]=='한글펌프'))
            check('analysis Excel preserves selection metadata',book['Rule Selection']['B1'].value=='VERIFYLINE')
        check('missing equipment rule strict mode rejected',request('/projects/'+pid+'/analysis?carrier=VERIFYLINE&effectiveDate=2026-06-01&allowDefaultFallback=false',token=a)[0]==409)
        status,plan=request('/projects/'+pid+'/plans'+query,'POST',token=a)
        if status!=200:print('Plan response:',status,plan)
        check('selected carrier plan saved',status==200);planid=plan['id']
        with db.cursor() as cur:
            cur.execute('select input_snapshot,result from oog_clp_scenario where id=%s',(planid,));saved_input,saved_result=cur.fetchone()
            check('selected rule stored in immutable input',saved_input['ruleSelection']['carrier']=='VERIFYLINE' and saved_input['ruleSelection']['effectiveDate']=='2026-06-01')
        pdfpath='/projects/'+pid+'/plans/'+planid+'/export/pdf'
        check('PDF anonymous denied',request(pdfpath)[0]==401)
        check('PDF foreign company denied',request(pdfpath,token=b)[0]==404)
        status,pdf=request(pdfpath,token=a);check('saved CLP PDF download',status==200 and pdf.startswith(b'%PDF'))
        qa=ROOT/'docs'/'export-qa';qa.mkdir(exist_ok=True)
        (qa/'verified-clp.pdf').write_bytes(pdf)
        with fitz.open(stream=pdf,filetype='pdf') as document:
            text='\n'.join(page.get_text() for page in document)
            check('PDF Korean text preserved','한글펌프' in text)
            check('PDF status and both scenarios','PRE_PLAN' in text or 'PRE-PLAN' in text)
            check('PDF unassigned included','UNASSIGNED' in text)
            check('PDF all scenarios rendered','MIN_EQUIPMENT' in text and 'EASY_HANDLING' in text)
            pages=len(document)
            for i,page in enumerate(document):page.get_pixmap(matrix=fitz.Matrix(1.2,1.2)).save(qa/f'live-page-{i+1}.png')
        status,xlsx=request('/projects/'+pid+'/plans/'+planid+'/export',token=a)
        with closing(openpyxl.load_workbook(io.BytesIO(xlsx))) as book:
            check('carrier snapshot Excel remains available',status==200 and len(book.sheetnames)==11)
            values=[cell.value for row in book['Summary'] for cell in row]
            check('CLP Excel preserves saved carrier metadata','VERIFYLINE' in values and '2026-06-01' in values)
        with db.cursor() as cur:
            cur.execute('select input_snapshot,result from oog_clp_scenario where id=%s',(planid,))
            check('exports leave saved plan immutable',cur.fetchone()==(saved_input,saved_result))
            cur.execute("select count(*) filter(where tablename like 'core_%%'),count(*) filter(where tablename like 'oog_%%'),count(*) from pg_tables where schemaname='public'")
            check('eight common sixteen business twenty-five total tables',cur.fetchone()==(8,16,25))
            cur.execute('select version from flyway_schema_history where success order by installed_rank desc limit 1')
            check('Flyway V10 applied',cur.fetchone()[0]=='10')
    finally:
        # Table lock and trigger change remain in this single transaction; other sessions
        # never see disabled immutability. Only this run's tenant IDs are deleted.
        db.execute('alter table oog_carrier_rule_version disable trigger oog_carrier_rule_immutable')
        for tid,_ in tenants:
            db.execute('delete from core_audit_history where tenant_id=%s',(tid,))
            for table in ['oog_clp_scenario','oog_packing_list','oog_cargo']:
                db.execute(f'delete from {table} where project_id in(select id from oog_project where tenant_id=%s)',(tid,))
            db.execute('delete from oog_project where tenant_id=%s',(tid,))
            db.execute('delete from oog_carrier_rule_version where tenant_id=%s',(tid,))
            for table in ['core_login_history','core_session','core_user_account','core_role_menu','core_menu','core_role']:
                db.execute(f'delete from {table} where tenant_id=%s',(tid,))
            db.execute('delete from core_tenant where tenant_id=%s',(tid,))
        db.execute('alter table oog_carrier_rule_version enable trigger oog_carrier_rule_immutable')
        db.execute("delete from core_audit_history where action='EQUIPMENT_UPDATED' and detail->>'equipmentCode'=%s",(fixture_code,))
        db.execute("delete from core_audit_history where action='EQUIPMENT_CREATED' and detail->>'code'=%s",(fixture_code,))
        db.execute('delete from oog_carrier_rule where equipment_code=%s',(fixture_code,))
        db.execute('delete from oog_equipment where code=%s',(fixture_code,))
        for jti in sessions:db.execute('update core_session set revoked_at=GREATEST(now(),created_at) where access_jti=%s',(jti,))
        db.commit()
total=failures=errors=0
for path in (ROOT/'backend/build/test-results/test').glob('TEST-*.xml'):
    node=ET.parse(path).getroot();total+=int(node.get('tests',0));failures+=int(node.get('failures',0));errors+=int(node.get('errors',0))
assert failures==errors==0
report={'date':'2026-10-13','migration':'V10','database':'oog','javaTests':total,'apiChecks':len(checks),'checks':checks,'pdfPages':pages,'fixturesRemoved':True,'uiWorkflow':'not tested','tables':{'core':8,'business':16,'flyway':1,'total':25}}
(ROOT/'docs/verification-master-pdf.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({k:report[k] for k in ['apiChecks','javaTests','pdfPages','tables','fixturesRemoved']}))
