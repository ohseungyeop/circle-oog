"""Build the maintained Korean user manual from the editable JSON source."""
import json
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[1]
source = json.loads((ROOT / 'docs/user-manual.json').read_text(encoding='utf-8'))
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Cm(21), Cm(29.7)
section.top_margin = section.bottom_margin = Cm(1.8)
section.left_margin = section.right_margin = Cm(1.9)
section.footer_distance = Cm(0.8)
for name, size in [('Normal', 10), ('Title', 23), ('Subtitle', 11), ('Heading 1', 17), ('Heading 2', 12), ('Caption', 9), ('List Number',10), ('List Bullet',10)]:
    style = doc.styles[name]
    style.font.name = 'Malgun Gothic'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0,0,0)
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Malgun Gothic')
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.15
doc.styles['Heading 1'].paragraph_format.space_before = Pt(8)
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
doc.styles['Subtitle'].font.italic = False
for border in doc.styles.element.xpath('.//w:pBdr'):
    border.getparent().remove(border)
doc.core_properties.author = ''
doc.core_properties.last_modified_by = ''
doc.core_properties.title = source['title']
doc.core_properties.subject = 'Circle OOG 사용자 업무 절차'
doc.core_properties.comments = ''
footer=section.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run(f"Circle OOG  |  {source['version']}  |  ")
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def paragraph(text, style=None):
    p=doc.add_paragraph(text,style)
    p.paragraph_format.widow_control=True
    return p

def table(headers,rows,widths):
    tbl=doc.add_table(rows=1,cols=len(headers))
    tbl.alignment=WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit=False
    for cell,width in zip(tbl.columns,widths):cell.width=Cm(width)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        item=OxmlElement('w:'+side)
        for key,value in [('val','single'),('sz','4'),('color','D9D9D9')]:item.set(qn('w:'+key),value)
        borders.append(item)
    tbl._tbl.tblPr.append(borders)
    for index,values in enumerate([headers]+rows):
        row=tbl.rows[0] if index==0 else tbl.add_row()
        if index==0:
            repeat=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(repeat)
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for i,(cell,text) in enumerate(zip(row.cells,values)):
            cell.width=Cm(widths[i]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
            run=p.add_run(text);run.font.size=Pt(9)
            if index==0:run.bold=True
            tcPr=cell._tc.get_or_add_tcPr()
            cell_borders=OxmlElement('w:tcBorders')
            for side in ['top','left','bottom','right']:
                item=OxmlElement('w:'+side)
                for border_key,border_value in [('val','single'),('sz','4'),('color','D9D9D9')]:item.set(qn('w:'+border_key),border_value)
                cell_borders.append(item)
            tcPr.append(cell_borders)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF6' if index==0 else 'FFFFFF');tcPr.append(shade)
            margins=OxmlElement('w:tcMar')
            for side in ['top','left','bottom','right']:
                item=OxmlElement('w:'+side);item.set(qn('w:w'),'90');item.set(qn('w:type'),'dxa');margins.append(item)
            tcPr.append(margins)
    paragraph('')

for index,page in enumerate(source['pages']):
    if index==0:
        paragraph(source['title'],'Title')
        paragraph(f"사용자 업무 안내  {source['version']}  {source['updated']}",'Subtitle')
    heading=paragraph(page['title'],'Heading 1')
    heading.paragraph_format.page_break_before=bool(index)
    for block in page['blocks']:
        if 'heading' in block:paragraph(block['heading'],'Heading 2')
        if 'text' in block:paragraph(block['text'])
        if 'steps' in block:
            for number,text in enumerate(block['steps'],1):paragraph(f'{number}. {text}')
        if 'bullets' in block:
            for text in block['bullets']:paragraph(text,'List Bullet')
        if 'table' in block:table(**block['table'])
        if 'image' in block:
            path=ROOT/block['image']
            if not path.exists():raise FileNotFoundError(path)
            p=paragraph('');p.add_run().add_picture(str(path),width=Cm(17.2))
            paragraph(block['caption'],'Caption')

output=ROOT/'docs/Circle-OOG-User-Manual.docx'
doc.save(output)
print(f'Created {output.name}; sections={len(source["pages"])}; version={source["version"]}')
