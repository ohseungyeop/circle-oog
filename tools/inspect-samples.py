import json
from pathlib import Path
import openpyxl

source=Path(r'D:\01. 오승엽\OOG')
result=[]
for path in source.glob('*.xlsx'):
    with path.open('rb') as stream:
        book=openpyxl.load_workbook(stream,read_only=True,data_only=True)
        sheets=[]
        for sheet in book.worksheets:
            rows=[]
            for row in sheet.iter_rows(min_row=1,max_row=min(12,sheet.max_row or 12),max_col=min(20,sheet.max_column or 20),values_only=True):
                rows.append([str(v) if v is not None else '' for v in row])
            sheets.append({'name':sheet.title,'rows':sheet.max_row,'columns':sheet.max_column,'preview':rows})
        result.append({'file':path.name,'sheets':sheets})
        book.close()
target=Path(__file__).resolve().parents[1]/'docs'/'sample-analysis.json'
target.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
print(f'Inspected {len(result)} Excel samples; report: {target}')
