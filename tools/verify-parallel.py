import base64
import io
import json
import time
import sys
import uuid
import urllib.request
import urllib.error
import zipfile
from xml.sax.saxutils import escape
from pathlib import Path
from urllib.parse import urlparse
import psycopg
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
config=dict(line.split('=',1) for line in (ROOT/'backend/.env').read_text(encoding='utf-8-sig').splitlines() if '=' in line)
db_url=urlparse(config['DB_URL'].removeprefix('jdbc:'));assert db_url.path=='/oog'
token=''
checks=[];project_ids=[]

def check(name,ok):
    assert ok,name
    checks.append(name)

def request(path,method='GET',data=None,content_type='application/json',auth=True):
    headers={'Content-Type':content_type}
    if auth:headers['Authorization']='Bearer '+token
    body=json.dumps(data).encode() if isinstance(data,(dict,list)) else data
    try:
        with urllib.request.urlopen(urllib.request.Request('http://127.0.0.1:8083/api'+path,data=body,headers=headers,method=method),timeout=45) as response:
            raw=response.read()
            return response.status,json.loads(raw).get('data') if response.headers.get_content_type().endswith('json') else raw
    except urllib.error.HTTPError as e:return e.code,e.read()

def fixture():
    rows=[['Item No','Description','Quantity','Original Dimension','','','Box Dimension','','','Net Weight','Gross Weight'],
          ['','','','Length','Width','Height','Length','Width','Height','',''],
          ['A','=SUM(1,2)',2,.9,.8,.7,1,1,1,1.6,2],
          ['B','Pump',1,1.5,.8,.9,2,1,1,1,1.2],
          ['X','Oversize',1,29,2,3,30,3,4,1,2]]
    sheet=['<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData>']
    for r,values in enumerate(rows,1):
        sheet.append(f'<row r="{r}">')
        for c,v in enumerate(values):
            ref=chr(65+c)+str(r)
            sheet.append(f'<c r="{ref}" t="inlineStr"><is><t>{escape(str(v))}</t></is></c>')
        sheet.append('</row>')
    sheet.append('</sheetData><mergeCells count="2"><mergeCell ref="D1:F1"/><mergeCell ref="G1:I1"/></mergeCells></worksheet>')
    out=io.BytesIO()
    with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
        z.writestr('[Content_Types].xml','<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/><Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/></Types>')
        z.writestr('_rels/.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/></Relationships>')
        z.writestr('xl/workbook.xml','<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><sheets><sheet name="PL" sheetId="1" r:id="rId1"/></sheets></workbook>')
        z.writestr('xl/_rels/workbook.xml.rels','<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/></Relationships>')
        z.writestr('xl/worksheets/sheet1.xml',''.join(sheet))
    return out.getvalue()

def upload(path,fields,raw):
    boundary='oog-'+uuid.uuid4().hex
    chunks=[f'--{boundary}\r\nContent-Disposition: form-data; name="{key}"\r\n\r\n{value}\r\n'.encode() for key,value in fields.items()]
    chunks += [f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="parallel.xlsx"\r\nContent-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n'.encode(),raw,f'\r\n--{boundary}--\r\n'.encode()]
    return request(path,'POST',b''.join(chunks),'multipart/form-data; boundary='+boundary)

for i in range(45):
    try:
        status,login=request('/auth/login','POST',{'tenantCode':'OOG','username':'admin','password':config['OOG_ADMIN_PASSWORD']},auth=False)
        if status==200:
            token=login['accessToken'];break
    except urllib.error.URLError:pass
    time.sleep(1)
else:raise RuntimeError('OOG API did not become ready')

if '--status-only' in sys.argv:
    status,payload=request('/development/status')
    canonical=json.loads((ROOT/'docs/feature-status.json').read_text(encoding='utf-8'))
    expected=json.loads((ROOT/'docs/verification-parallel.json').read_text(encoding='utf-8'))['numericSummary']
    assert status==200 and payload['summary']==expected
    assert payload['items']==canonical['items'] and payload['version']==canonical['version']
    with psycopg.connect(host=db_url.hostname,port=db_url.port,dbname='oog',user=config['DB_USERNAME'],password=config['DB_PASSWORD'],connect_timeout=10) as cleanup_db:
        claims=json.loads(base64.urlsafe_b64decode(token.split('.')[1]+'=='))
        cleanup_db.execute('update core_session set revoked_at=GREATEST(now(),created_at) where access_jti=%s',(claims['jti'],))
    print(json.dumps({'serverSummary':payload['summary'],'canonicalMatches':True}))
    raise SystemExit(0)

with psycopg.connect(host=db_url.hostname,port=db_url.port,dbname='oog',user=config['DB_USERNAME'],password=config['DB_PASSWORD'],connect_timeout=10) as db:
    try:
        check('new history endpoint requires authentication',request('/projects/'+str(uuid.uuid4())+'/history',auth=False)[0]==401)
        for suffix in ['MAIN','OTHER']:
            status,p=request('/projects','POST',{'name':'VERIFY-PARALLEL-'+suffix+'-'+uuid.uuid4().hex,'customer':'Verification'})
            check('project created '+suffix,status==200);project_ids.append(p['id'])
        pid,other=project_ids
        raw=fixture();status,preview=upload('/imports/preview',{'header':0,'headerRows':2},raw)
        check('two header rows previewed',status==200 and len(preview['combinedHeaders'])==11)
        mapping=preview['suggestedMapping']
        check('packed and original columns distinguished',mapping['length']==6 and mapping['originalLength']==3 and mapping['netWeight']==9 and mapping['grossWeight']==10)
        fields={'mapping':json.dumps(mapping),'header':0,'headerRows':2,'lengthUnit':'m','weightUnit':'ton','weightBasis':'ROW_TOTAL'}
        check('malformed mapping is HTTP 400',upload('/projects/'+pid+'/imports',{**fields,'mapping':'{bad'},raw)[0]==400)
        check('invalid weight basis rejected',upload('/projects/'+pid+'/imports',{**fields,'weightBasis':'AUTO'},raw)[0]==400)
        status,imported=upload('/projects/'+pid+'/imports',fields,raw)
        check('row totals imported atomically',status==200 and imported['imported']==3)
        status,cargos=request('/projects/'+pid+'/cargos');a=next(c for c in cargos if c['itemNo']=='A')
        check('meter dimensions stored as mm',a['length']==1000 and a['width']==1000 and a['height']==1000)
        check('original dimensions retained independently',a['originalLength']==900 and a['originalWidth']==800 and a['originalHeight']==700)
        check('ton row gross/net totals divided by quantity',a['quantity']==2 and a['grossWeight']==1000 and a['netWeight']==800)
        status,duplicate=upload('/projects/'+pid+'/imports',fields,raw)
        check('duplicate upload creates no rows',status==200 and duplicate['imported']==0)
        status,history=request('/projects/'+pid+'/history?limit=50')
        check('uploads returned as structured mapping',status==200 and len(history['uploads'])==1 and history['uploads'][0]['mapping']['weightBasis']=='ROW_TOTAL')
        upload_id=history['uploads'][0]['id']
        check('download is exact original bytes',request('/projects/'+pid+'/imports/'+upload_id+'/file')==(200,raw))
        check('foreign project original rejected',request('/projects/'+other+'/imports/'+upload_id+'/file')[0]==404)
        check('history pagination limit enforced',request('/projects/'+pid+'/history?limit=101')[0]==400)
        status,plan=request('/projects/'+pid+'/plans','POST');check('CLP created',status==200)
        plan_id=plan['id'];source=plan['result'];s=source['scenarios'][0];box=s['containers'][0]
        assignment=next(v for v in box['assignments'] if v['cargoId']==a['id'] and v['piece']==1)
        check('oversize stays unassigned',len(s['unassigned'])==1 and s['unassigned'][0]['itemNo']=='X')
        change={'cargoId':a['id'],'piece':1,'targetContainerNumber':box['number'],'x':5000,'y':assignment['y'],'locked':True}
        status,revision=request('/projects/'+pid+'/plans/'+plan_id+'/revisions','POST',{'scenarioIndex':0,'changes':[change]})
        check('move plus lock creates a new plan',status==200 and revision['id']!=plan_id and revision['sourcePlanId']==plan_id)
        locked_id=revision['id'];locked_source=revision['result']
        check('unselected scenario unchanged',locked_source['scenarios'][1]==source['scenarios'][1])
        locked_assignment=next(v for b in locked_source['scenarios'][0]['containers'] for v in b['assignments'] if v['cargoId']==a['id'] and v['piece']==1)
        check('new coordinates and lock retained',locked_assignment['x']==5000 and locked_assignment['locked'])
        invalid={**change,'x':6000};invalid.pop('locked')
        check('locked move rejected without explicit unlock',request('/projects/'+pid+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[invalid]})[0]==400)
        check('overlap rejected',request('/projects/'+pid+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[{**change,'x':1000,'locked':False}]})[0]==400)
        check('equipment length limit enforced',request('/projects/'+pid+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[{**change,'x':999999,'locked':False}]})[0]==400)
        check('rotation permission enforced',request('/projects/'+pid+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[{**change,'rotated':True,'locked':False}]})[0]==400)
        check('foreign project revision rejected',request('/projects/'+other+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[change]})[0]==404)
        status,unlocked=request('/projects/'+pid+'/plans/'+locked_id+'/revisions','POST',{'scenarioIndex':0,'changes':[{**change,'x':6000,'locked':False}]})
        check('explicit unlock and move accepted',status==200 and unlocked['id']!=locked_id)
        status,history=request('/projects/'+pid+'/history?size=50')
        check('source plan remains immutable',next(p for p in history['plans'] if p['id']==plan_id)['result']==source)
        check('invalid revisions saved no rows',len(history['plans'])==3)
        check('revision audit contains source lineage',sum(e['action']=='CLP_REVISED' for e in history['audits'])==2 and all(isinstance(e['detail'],dict) for e in history['audits']))
        status,export=request('/projects/'+pid+'/plans/'+unlocked['id']+'/export');check('CLP Excel exported',status==200)
        with io.BytesIO(export) as stream:
            book=openpyxl.load_workbook(stream,read_only=False)
            check('all 11 CLP sheets present',set(book.sheetnames)=={'Summary','Assignments','Containers','Input Cargos','Equipment','Rules','Unassigned','Warnings','Top','Side','Front'})
            check('assignments from all saved scenarios exported',book['Assignments'].max_row==7)
            check('unassigned from both scenarios visible',book['Unassigned'].max_row==3)
            rows=list(book['Assignments'].values);headers=rows[0];idx=headers.index('scenarioIndex')
            check('scenario identity retained in Excel',set(row[idx] for row in rows[1:])=={0,1})
            check('all three diagrams contain colored cargo',all(any(cell.fill.patternType=='solid' for row in book[name] for cell in row) for name in ['Top','Side','Front']))
            check('customer formula text stays text',any(cell.value=='=SUM(1,2)' and cell.data_type=='s' for row in book['Input Cargos'] for cell in row))
            book.close()
        check('foreign project plan export rejected',request('/projects/'+other+'/plans/'+unlocked['id']+'/export')[0]==404)
        status,stats=request('/development/status')
        check('development status counts add up',status==200 and stats['summary']['total']==len(json.loads((ROOT/'docs/feature-status.json').read_text(encoding='utf-8'))['items']) and stats['summary']['completed']+stats['summary']['incomplete']==stats['summary']['total'] and stats['summary']['verified']+stats['summary']['verificationPending']==stats['summary']['total'])
    finally:
        with db.cursor() as cursor:
            for pid in project_ids:
                for table in ['core_audit_history','oog_clp_scenario','oog_packing_list','oog_cargo']:
                    cursor.execute(f'delete from {table} where project_id=%s',(pid,))
                cursor.execute('delete from oog_project where id=%s',(pid,))
        db.commit()
        import shutil
        for pid in project_ids:
            folder=ROOT/'data'/'projects'/pid
            if folder.exists() and folder.resolve().is_relative_to((ROOT/'data').resolve()):shutil.rmtree(folder)
with psycopg.connect(host=db_url.hostname,port=db_url.port,dbname='oog',user=config['DB_USERNAME'],password=config['DB_PASSWORD'],connect_timeout=10) as cleanup_db:
    claims=json.loads(base64.urlsafe_b64decode(token.split('.')[1]+'=='))
    cleanup_db.execute('update core_session set revoked_at=GREATEST(now(),created_at) where access_jti=%s',(claims['jti'],))
report={'date':'2026-10-13','database':'oog','migration':'V10','tables':{'core':8,'business':16,'flyway':1,'total':25},'apiChecks':len(checks),'checks':checks,'testProjectsRemoved':True,'uiWorkflow':'not tested','javaTests':None,'frontendBuild':'passed','scope':['P03','I04','L04','O03','H01','numeric development status']}
(ROOT/'docs/verification-parallel.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'apiChecks':len(checks),'database':'oog','fixturesRemoved':True}))
